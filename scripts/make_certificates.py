# -*- coding: utf-8 -*-
"""Generate a black & white, elegantly formatted Word document of family
certificates for a 65th birthday celebration.

Layout:
  - Page 1: title (full page, double frame)
  - Page 2: instructions (full page, thin frame)
  - Pages 3..: two certificates per page (half-page each), no blank pages.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Palatino Linotype"
OUT = r"D:\pi\Семейные_сертификаты_65_лет.docx"

# --- Content ---------------------------------------------------------------
CERTIFICATES = [
    {"n": 1, "title": "«Генеральная уборка»",
     "body": "Владелец данного документа имеет право вызвать указанного "
             "гражданина на генеральную уборку в квартире именинницы. "
             "Включает: пыль, полы, окна и фразу «Сядь, отдохни, я сам(а)»."},
    {"n": 2, "title": "«Шеф-повар на вечер»",
     "body": "Указанное лицо обязано приготовить для именинницы ужин из трёх "
             "блюд. Меню утверждает юбиляр. Жалобы на тему «я не умею готовить» "
             "не принимаются. Допускается один звонок за подсказкой."},
    {"n": 3, "title": "«Личный водитель»",
     "body": "Именинница имеет право в любой момент сказать: «Мне нужно в "
             "поликлинику / на рынок / к подруге», — и указанный гражданин "
             "обязан бросить всё, сесть за руль и отвезти. Музыка в машине — "
             "по выбору юбиляра."},
    {"n": 4, "title": "«Дачный спецназ на 3 часа»",
     "body": "Указанный гражданин по первому зову именинницы выезжает на дачу "
             "и в течение трёх часов выполняет команды: «Копай тут», «Полей "
             "там», «А вот тут ещё прополи». Протесты в духе «У меня спина» "
             "не засчитываются."},
    {"n": 5, "title": "«Вечер без телефона»",
     "body": "Указанное лицо обязано провести с именинницей два часа без "
             "единого взгляда в телефон. Разрешается: пить чай, слушать "
             "истории, играть в лото, смотреть альбомы. За каждый взгляд в "
             "экран — штраф: комплимент юбиляру."},
    {"n": 6, "title": "«Поход в бассейн»",
     "body": "Указанный гражданин обязан сопроводить именинницу в бассейн: "
             "донести сумку, занять удобную дорожку и вовремя подбодрить "
             "фразой «Ты ещё ого-го!». Отговорка «Я плавки забыл» не "
             "засчитывается — запасные выдаёт юбиляр."},
    {"n": 7, "title": "«Караоке по заявке»",
     "body": "Именинница называет песню, и указанное лицо обязано спеть её "
             "на следующем семейном застолье. Неважно, есть ли слух. "
             "Неважно, знают ли слова. Главное — душа и громкость."},
    {"n": 8, "title": "«Техподдержка на дому»",
     "body": "Указанный гражданин становится личным IT-специалистом "
             "именинницы. Обязан: настроить телефон, объяснить, куда делась "
             "кнопка, почему не работает мессенджер и что такое «эта ваша "
             "облака». Без вздохов и закатывания глаз."},
    {"n": 9, "title": "«Поход в театр»",
     "body": "Указанное лицо обязано сопроводить именинницу в театр: купить "
             "программку, сдать пальто в гардероб и не засыпать во втором "
             "действии. Шуршание конфетными обёртками допускается, но тихо "
             "и не в лирических сценах."},
    {"n": 10, "title": "«Семейный фотограф»",
     "body": "Указанный гражданин назначается личным фотографом юбиляра на "
             "любом семейном событии. Обязан сделать минимум 30 кадров, из "
             "которых именинница выберет два удачных. Фраза «Ну хватит уже "
             "фоткаться» запрещена."},
    {"n": 11, "title": "«Концерт органной музыки»",
     "body": "Указанный гражданин сопровождает именинницу на концерт органной "
             "музыки. Обязан сидеть с возвышенным лицом, не шептать «А долго "
             "ещё?» и аплодировать строго в нужных местах. Во время "
             "исполнения — абсолютная тишина."},
    {"n": 12, "title": "«Помощник по растениям»",
     "body": "Указанный гражданин обязуется помочь имениннице с пересадкой "
             "цветов, протиркой листьев или поливом огорода. Запрещается "
             "задавать вопросы в духе «Зачем тебе столько герани?» или "
             "«Оно же всё равно засохнет»."},
    {"n": 13, "title": "«Сопровождающий на прогулке»",
     "body": "Указанный гражданин обязан выделить полдня, чтобы неспешно "
             "прогуляться с именинницей в парке или по любимым местам, идя "
             "строго в её темпе и активно поддерживая беседу."},
    {"n": 14, "title": "«Танец на следующем празднике»",
     "body": "На следующем семейном застолье указанное лицо обязано "
             "станцевать с именинницей медленный танец. Песню выбирает "
             "юбиляр. Отговорки в стиле «Я не умею танцевать» отклоняются "
             "комиссией в составе всей семьи."},
    {"n": 15, "title": "«Помощник по заготовкам»",
     "body": "В сезон заготовок указанный гражданин обязуется чистить, "
             "резать, мешать и закатывать банки под строгим руководством "
             "именинницы. Фразы «Зачем так много» и «Мы это всё не съедим» "
             "запрещены. Разрешается честно и много пробовать."},
]

# --- Low-level helpers -----------------------------------------------------
BLACK = (0, 0, 0)
_ALIGN = {
    'center': WD_ALIGN_PARAGRAPH.CENTER,
    'left': WD_ALIGN_PARAGRAPH.LEFT,
    'right': WD_ALIGN_PARAGRAPH.RIGHT,
    'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def style_run(run, name=FONT, size=11, bold=False, italic=False,
              color=BLACK, spacing=None, small_caps=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(attr), name)
    if spacing is not None:
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:val'), str(spacing))
        rpr.append(sp)
    if small_caps:
        rpr.append(OxmlElement('w:smallCaps'))


def add_para(parent, text, size=11, bold=False, italic=False, align='center',
             name=FONT, spacing=None, small_caps=False, space_before=0,
             space_after=0, color=BLACK, line_spacing=None):
    p = parent.add_paragraph()
    p.alignment = _ALIGN[align]
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.first_line_indent = Cm(0)
    if line_spacing is not None:
        pf.line_spacing = line_spacing
    r = p.add_run(text)
    style_run(r, name=name, size=size, bold=bold, italic=italic, color=color,
              spacing=spacing, small_caps=small_caps)
    return p


def add_ornament(parent, text="\u2726", size=12, space_before=0, space_after=0):
    return add_para(parent, text, size=size, align='center',
                    space_before=space_before, space_after=space_after)


# --- Table / cell frame helpers -------------------------------------------
def _set_cell_borders(cell, sz=8, val="double", color="000000"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), val)
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)
        borders.append(e)
    tcPr.append(borders)


def _clear_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'nil')
        borders.append(e)
    tcPr.append(borders)


def _set_cell_margins(cell, top=350, bottom=350, left=600, right=600):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for name, val in (('top', top), ('left', left),
                      ('bottom', bottom), ('right', right)):
        node = OxmlElement(f'w:{name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        mar.append(node)
    tcPr.append(mar)


def _full_width_table(doc):
    table = doc.add_table(rows=0, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')
    tblPr.append(tblW)
    return table


def _add_row(table, height_cm, rule="atLeast", cant_split=True):
    row = table.add_row()
    trPr = row._tr.get_or_add_trPr()
    if cant_split:
        trPr.append(OxmlElement('w:cantSplit'))
    h = OxmlElement('w:trHeight')
    h.set(qn('w:val'), str(int(height_cm * 567)))
    h.set(qn('w:hRule'), rule)
    trPr.append(h)
    return row


def _styled_cell(row, val="double", sz=8, margins=(200, 200, 500, 500)):
    cell = row.cells[0]
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _set_cell_borders(cell, sz=sz, val=val)
    _set_cell_margins(cell, *margins)
    return cell


def clean_first_empty(cell):
    p0 = cell.paragraphs[0]
    if not p0.text and not p0.runs:
        p0._element.getparent().remove(p0._element)


def page_break(doc):
    """Invisible 1pt paragraph that forces the next content onto a new page.
    Sized so it never overflows and never creates a blank page."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.page_break_before = True
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(1)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    r = p.add_run('')
    style_run(r, size=1)
    return p


def setup_page(doc):
    s = doc.sections[0]
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.top_margin = Cm(2.0)
    s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(2.4)
    s.right_margin = Cm(2.4)


def set_default_font(doc):
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(11)
    # Zero out spacing so the implicit paragraph Word inserts after every
    # table stays minimal and never pushes content onto a blank page.
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(attr), FONT)


# --- Page builders ---------------------------------------------------------
# NOTE: for hRule="atLeast", Word renders row = trHeight + cell top/bottom
# margins + border. So the values below are CONTENT heights; the visible
# frame is slightly taller. Usable page height ~ 25.7 cm.
FULL_H = 24.2     # full-page content height (frame ~25.0 cm)
HALF_H = 11.2     # half-page certificate content height (frame ~12.0 cm)
GAP_H = 0.45      # spacer row height between the two certificates (exact)


def fill_title_cell(cell):
    add_ornament(cell, "\u2726  \u2726  \u2726", size=18, space_after=20)
    add_para(cell, "СЕКРЕТНЫЕ", size=16, bold=True, small_caps=True,
             spacing=80, space_after=8)
    add_para(cell, "СЕМЕЙНЫЕ", size=40, bold=True, spacing=60, space_after=0)
    add_para(cell, "СЕРТИФИКАТЫ", size=40, bold=True, spacing=60, space_after=14)
    add_ornament(cell, "\u2726  \u2726  \u2726", size=16, space_after=16)
    add_para(cell, "Юбилей · 65 лет", size=28, italic=True, space_after=28)
    add_para(cell, "Подарочно-развлекательное издание для всей семьи",
             size=13, italic=True, space_before=24, space_after=10)
    add_ornament(cell, "\u2726", size=14, space_before=10)
    clean_first_empty(cell)


def fill_instructions_cell(cell):
    add_para(cell, "Как пользоваться", size=24, bold=True, space_after=12)
    add_ornament(cell, "\u2726  \u2726  \u2726", size=13, space_after=18)
    add_para(cell,
             "Впишите в пустое поле имя того члена семьи, который, по "
             "вашему мнению, идеально подходит для исполнения этого "
             "сертификата.",
             size=15, italic=True, align='justify', line_spacing=1.4,
             space_after=30)
    add_ornament(cell, "\u2726", size=11, space_after=8)
    add_para(cell, "Одобрено семейным советом · Обжалованию не подлежит",
             size=12, italic=True, small_caps=True, spacing=30)
    clean_first_empty(cell)


def fill_certificate_cell(cell, cert):
    add_para(cell, f"СЕМЕЙНЫЙ СЕРТИФИКАТ № {cert['n']}", size=12, bold=True,
             small_caps=True, spacing=50, space_after=6)
    add_ornament(cell, "\u2726", size=10, space_after=8)
    add_para(cell, cert['title'], size=20, bold=True, italic=True,
             space_after=10)
    add_ornament(cell, "\u2726  \u2726  \u2726", size=9, space_after=10)
    add_para(cell, cert['body'], size=11, align='justify',
             line_spacing=1.3, space_after=14)
    add_para(cell, "Исполнитель: ______________________________",
             size=11, align='left', space_after=7)
    add_para(cell, "Подпись именинницы: ____________________",
             size=11, align='left', space_after=14)
    add_ornament(cell, "\u2726", size=9, space_after=3)
    add_para(cell, "Одобрено семейным советом · Обжалованию не подлежит",
             size=9.5, italic=True, small_caps=True, spacing=40)
    clean_first_empty(cell)


def build_full_page(doc, filler, val="double", sz=12,
                    margins=(200, 200, 700, 700)):
    table = _full_width_table(doc)
    row = _add_row(table, FULL_H)
    cell = _styled_cell(row, val=val, sz=sz, margins=margins)
    filler(cell)


def build_cert_pair_page(doc, cert_a, cert_b=None):
    table = _full_width_table(doc)
    # Top certificate.
    row_a = _add_row(table, HALF_H)
    cell_a = _styled_cell(row_a, val="double", sz=8)
    fill_certificate_cell(cell_a, cert_a)
    if cert_b is not None:
        # Spacer row (no borders, small fixed gap).
        row_gap = _add_row(table, GAP_H, rule="exact", cant_split=True)
        cell_gap = row_gap.cells[0]
        _clear_cell_borders(cell_gap)
        _set_cell_margins(cell_gap, 0, 0, 0, 0)
        gp = cell_gap.paragraphs[0]
        gp.paragraph_format.line_spacing = Pt(1)
        gp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        gp.paragraph_format.space_after = Pt(0)
        # Bottom certificate.
        row_b = _add_row(table, HALF_H)
        cell_b = _styled_cell(row_b, val="double", sz=8)
        fill_certificate_cell(cell_b, cert_b)


def build_document():
    doc = Document()
    setup_page(doc)
    set_default_font(doc)
    doc.core_properties.title = "Семейные сертификаты — 65 лет"

    # Page 1: title.
    build_full_page(doc, fill_title_cell, val="double", sz=12)

    # Page 2: instructions.
    page_break(doc)
    build_full_page(doc, fill_instructions_cell, val="single", sz=4)

    # Pages 3..: two certificates per page.
    i = 0
    while i < len(CERTIFICATES):
        page_break(doc)
        a = CERTIFICATES[i]
        b = CERTIFICATES[i + 1] if i + 1 < len(CERTIFICATES) else None
        build_cert_pair_page(doc, a, b)
        i += 2

    doc.save(OUT)
    print("Saved:", OUT)


if __name__ == "__main__":
    build_document()
